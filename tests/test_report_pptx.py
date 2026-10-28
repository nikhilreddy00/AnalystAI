"""Gap 4 — smoke test for the PowerPoint export rebuild.

Builds a deck from a realistic AnalysisSession fixture and asserts the
structural invariants:
- Output is a valid .pptx (ZIP magic).
- Slide count matches expected (title + KPI + summary + N charts + actions + closing).
- Title slide and closing slide have navy backgrounds.
- No slide carries the legacy indigo accent rectangle.
- All slide titles use Georgia.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from api.lib.report_pptx import NAVY, TITLE_FONT, build_pptx
from orchestrator.state import (
    ActionItem,
    AnalysisFindings,
    AnalysisSession,
    ChartConfig,
    ChartType,
    ExecutiveSummary,
    InsightReport,
    KPI,
)


@pytest.fixture
def fixture_session() -> AnalysisSession:
    """Realistic mini-session — enough to populate every slide kind."""
    ev = ExecutiveSummary(
        headline="Women survived at 4x the rate of men",
        paragraphs=[],
        findings=[
            "Female passengers survived at 74.2% vs 18.9% for males",
            "First class survival was 62.9% vs 24.2% for third class",
            "Cabin column is 40%+ missing across the dataset",
        ],
        action_items=[
            ActionItem(action="Model using gender, class, age, fare", expected_impact="primary predictors"),
            ActionItem(action="Drop cabin column from feature set", expected_impact="too sparse to impute reliably"),
        ],
    )
    report = InsightReport(
        executive_summary="Women survived at 74.2% vs 18.9% for men.",
        technical_findings=[],
        recommended_actions=[],
        anomaly_alerts=[],
        proactive_questions=[],
        methodology_notes="",
        executive_view=ev,
    )
    findings = AnalysisFindings(
        eda_summary={},
        correlations={},
        trends=[],
        segments=[],
        anomalies=[],
        statistical_tests=[],
        code_executed="",
        kpis=[
            KPI(label="Survival Rate", value="38.4%", context="of 891 passengers"),
            KPI(label="Avg Fare", value="$32.20", context="median $14.45"),
            KPI(label="Top Class", value="3rd", context="55.1% of passengers"),
        ],
    )
    chart = ChartConfig(
        id="c1",
        chart_type=ChartType.BAR,
        title="Survival rate by sex",
        data=[{"sex": "female", "rate": 74.2}, {"sex": "male", "rate": 18.9}],
        x_key="sex",
        y_keys=["rate"],
        colors=["#1D9E75", "#A32D2D"],
        value_format="percent",
        source_finding="Female passengers survived at nearly 4x the rate of males.",
    )
    return AnalysisSession(
        session_id="t",
        user_id="u",
        raw_query="why did survival differ by class?",
        dataset_name="Titanic",
        chart_configs=[chart],
        analysis_findings=findings,
        insight_report=report,
    )


def test_build_pptx_returns_valid_zip(fixture_session):
    data = build_pptx(fixture_session)
    assert isinstance(data, bytes) and len(data) > 0
    assert data[:2] == b"PK"  # ZIP magic


def test_build_pptx_slide_count(fixture_session):
    """Title + KPI + summary + 1 chart + actions + closing = 6 slides."""
    data = build_pptx(fixture_session)
    prs = Presentation(BytesIO(data))
    assert len(prs.slides) == 6


def _slide_bg_rgb(slide) -> RGBColor | None:
    fill = slide.background.fill
    try:
        return fill.fore_color.rgb
    except Exception:
        return None


def test_title_and_closing_have_navy_background(fixture_session):
    data = build_pptx(fixture_session)
    prs = Presentation(BytesIO(data))
    assert _slide_bg_rgb(prs.slides[0]) == NAVY
    assert _slide_bg_rgb(prs.slides[-1]) == NAVY


def test_content_slides_have_non_navy_background(fixture_session):
    """Slides between title and closing should NOT be navy."""
    data = build_pptx(fixture_session)
    prs = Presentation(BytesIO(data))
    n = len(prs.slides)
    for i in range(1, n - 1):
        assert _slide_bg_rgb(prs.slides[i]) != NAVY


def test_no_legacy_accent_rectangle(fixture_session):
    """The pre-Gap-4 deck drew a tiny indigo rectangle (0.6"x0.05") below
    every slide title. The rewrite removes it. Verify no slide contains a
    non-text shape that small (excluding KPI shapes which are 1.8" tall)."""
    data = build_pptx(fixture_session)
    prs = Presentation(BytesIO(data))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                continue
            # 0.1" = 91440 EMUs
            if shape.height is not None and shape.height <= 91440:
                pytest.fail(f"suspicious tiny shape detected: height={shape.height} EMUs")


def test_slide_titles_use_georgia(fixture_session):
    """Spot-check: the title slide's main title run should have font name Georgia."""
    data = build_pptx(fixture_session)
    prs = Presentation(BytesIO(data))
    title_slide = prs.slides[0]
    fonts_seen = []
    for shape in title_slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if run.font.name:
                    fonts_seen.append(run.font.name)
    assert TITLE_FONT in fonts_seen, f"Expected Georgia on title slide; saw {fonts_seen}"


def test_grouped_bars_with_ci_and_empty_colours_render():
    """Phase 5: a Simpson-style grouped chart (2 series, CI whiskers, colors=[]) renders in the
    deck with distinct series colours instead of dropping the second series."""
    from api.lib import chart_render
    from orchestrator.state import ChartConfig, ChartType

    chart = ChartConfig(id="g", chart_type=ChartType.BAR, title="Conversion by department",
                        x_key="department", y_keys=["treatment", "control"], colors=[],
                        error_key="__ci", split_key="group", value_format="percent",
                        data=[{"department": "Easy", "treatment": 80.0, "treatment__ci": [5.0, 4.0],
                               "control": 70.0, "control__ci": [3.0, 3.0]},
                              {"department": "Hard", "treatment": 30.0, "treatment__ci": [3.0, 3.0],
                               "control": 20.0, "control__ci": [5.0, 6.0]}],
                        source_finding="group_bars")
    assert chart_render.render_chart_to_png(chart)[:4] == b"\x89PNG"
    assert len(set(chart_render._resolve_colors(chart, 2))) == 2
